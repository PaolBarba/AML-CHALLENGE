"""Excel snapshots of the authoritative challenge records."""
import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def workbook_bytes(store):
    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(name, headers, rows):
        page = workbook.create_sheet(name)
        page.append(headers)
        for values in rows:
            page.append([ILLEGAL_CHARACTERS_RE.sub('', v) if isinstance(v, str) else v for v in values])
        for row in page:
            for cell in row:
                # Participant text must never become an Excel formula.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                elif isinstance(cell.value, float):
                    cell.number_format = '0.0000'
        for cell in page[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='243B53')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        page.row_dimensions[1].height = 32
        page.freeze_panes = 'D2'
        page.auto_filter.ref = page.dimensions
        for index, header in enumerate(headers, 1):
            page.column_dimensions[get_column_letter(index)].width = min(38, max(18, len(header) + 2))

    rankings = []
    for split in ('train', 'validation', 'test'):
        for mode in ('quality', 'measured'):
            for item in store.leaderboard(split, measured=mode == 'measured'):
                rankings.append([split, mode, item['team'], item['run'], item['score'],
                                 item['completed'], item['total'],
                                 *[item['metrics'].get(k) for k in ('iou', 'editing', 'preservation', 'efficiency')],
                                 item['measured_total_time_ms']])
    sheet('Team rankings', ['Split', 'Mode', 'Team', 'Selected run', 'Score', 'Completed tasks', 'Total tasks',
                            'IoU', 'Editing', 'Preservation', 'Efficiency', 'Measured total time (ms)'], rankings)

    records = store.submissions('admin', admin=True)
    metric_names = sorted({key for record in records for key in record['metrics']})
    rows = []
    for record in records:
        task = store.tasks[record['task']]
        rows.append([record['id'], record['team'], record['run'], record['task'], record['split'],
                     task['task_type'], task['difficulty'],
                     datetime.fromtimestamp(record['created'], timezone.utc).isoformat(),
                     'measured' if record['measured_ms'] is not None else 'quality', record['measured_ms'],
                     *[record['metrics'].get(key) for key in metric_names]])
    sheet('Submissions', ['Submission ID', 'Team', 'Run', 'Task', 'Split', 'Task type', 'Difficulty',
                          'Submitted at (UTC)', 'Mode', 'Measured time (ms)', *metric_names], rows)
    output = io.BytesIO()
    batches = store.db.execute('SELECT id,team,source_hash,width,height,received FROM reference_batches WHERE received IS NOT NULL ORDER BY received').fetchall()
    sheet('Reference uploads', ['Batch ID', 'Team', 'Source SHA256', 'Width', 'Height', 'Received at (UTC)', 'Scoring status'],
          [[r['id'], r['team'], r['source_hash'], r['width'], r['height'],
            datetime.fromtimestamp(r['received'], timezone.utc).isoformat(), 'unconfigured'] for r in batches])
    workbook.save(output)
    workbook.close()
    return output.getvalue()
