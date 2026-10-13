import base64
import json
import math
import numpy as np
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from PIL import Image
from openpyxl import load_workbook
try:
    from fastapi.testclient import TestClient
    from challenge.api import create_app
except ModuleNotFoundError:
    TestClient = None

from challenge.reference_demo import initialize_reference
from challenge.reference_metrics import METRICS, pixel_metrics, ranking_key, ReferenceMetrics
from challenge.store import Store, image_bytes


def encoded(image):
    return base64.b64encode(image_bytes(image)).decode('ascii')


class ReferenceMetricTests(unittest.TestCase):
    def test_known_pixel_metrics_and_infinite_psnr(self):
        black = Image.new('RGB', (8, 8), 'black')
        identical = pixel_metrics(black, black)
        self.assertEqual(identical['mse'], 0)
        self.assertEqual(identical['ssim'], 1)
        self.assertIsNone(identical['psnr'])
        self.assertTrue(identical['psnr_infinite'])
        json.dumps(identical, allow_nan=False)
        different = pixel_metrics(black, Image.new('RGB', black.size, 'white'))
        self.assertEqual(different['mse'], 1)
        self.assertEqual(different['psnr'], 0)
        self.assertAlmostEqual(different['ssim'], .0001 / 1.0001)
        gray = pixel_metrics(black, Image.new('RGB', black.size, (128, 128, 128)))
        self.assertAlmostEqual(gray['mse'], (128 / 255) ** 2)
        self.assertAlmostEqual(gray['psnr'], 10 * math.log10(1 / gray['mse']))
        self.assertIsNone(pixel_metrics(Image.new('RGB', (1, 1)), Image.new('RGB', (1, 1)))['ssim'])
        with self.assertRaisesRegex(ValueError, 'matching dimensions'):
            pixel_metrics(black, Image.new('RGB', (4, 4)))

    def test_ssim_matches_direct_window_definition(self):
        rng = np.random.default_rng(7)
        a = rng.integers(0, 256, size=(9, 10, 3), dtype=np.uint8)
        b = rng.integers(0, 256, size=(9, 10, 3), dtype=np.uint8)
        values = []
        for y in range(3):
            for x in range(4):
                for channel in range(3):
                    left = a[y:y + 7, x:x + 7, channel].ravel() / 255
                    right = b[y:y + 7, x:x + 7, channel].ravel() / 255
                    covariance = np.cov(left, right, ddof=1)
                    luminance = (2 * left.mean() * right.mean() + .01 ** 2) / (left.mean() ** 2 + right.mean() ** 2 + .01 ** 2)
                    contrast = (2 * covariance[0, 1] + .03 ** 2) / (covariance[0, 0] + covariance[1, 1] + .03 ** 2)
                    values.append(luminance * contrast)
        actual = pixel_metrics(Image.fromarray(a), Image.fromarray(b))['ssim']
        self.assertAlmostEqual(actual, float(np.mean(values)), places=12)

    def test_models_missing_stay_blank_and_preservation_breaks_ties(self):
        image = Image.new('RGB', (8, 8))
        metrics = ReferenceMetrics({}).evaluate(image, image, {'target_prompt': 'Make it red'})
        for name in ('lpips', 'clip_similarity', 'structure_distance'):
            self.assertIsNone(metrics[name])
        a = dict(clip_similarity=.9, ssim=.3, lpips=.3, structure_distance=.3, mse=.3)
        b = dict(a, clip_similarity=.8, ssim=1)
        self.assertLess(ranking_key(a), ranking_key(b))
        self.assertLess(ranking_key(dict(a, ssim=.5)), ranking_key(a))
        self.assertLess(ranking_key(dict(a, lpips=.1)), ranking_key(a))
        self.assertLess(ranking_key(dict(a, structure_distance=.1)), ranking_key(a))
        with self.assertRaisesRegex(ValueError, 'model_version'):
            ReferenceMetrics({'clip_directory': '/nonexistent'})


@unittest.skipIf(TestClient is None, 'Install requirements-api.txt')
class ReferenceScoringTests(unittest.TestCase):
    def test_mask_free_upload_pixel_metrics_excel_and_persistence(self):
        with tempfile.TemporaryDirectory() as directory:
            root = initialize_reference(Path(directory) / 'demo', {'clip_directory': '/unavailable/clip', 'model_version': 'test'})
            config = json.loads((root / 'config.json').read_text())
            with TestClient(create_app(root)) as client:
                headers = {'Authorization': 'Bearer ' + config['teams']['team_demo']}
                self.assertEqual(client.get('/api/reference/leaderboard').status_code, 401)
                schema = client.get('/openapi.json').json()
                submission_schema = schema['paths']['/api/reference/submissions']['post']['requestBody']['content']['application/json']['schema']
                self.assertNotIn('predicted_mask', submission_schema['properties'])
                with patch('challenge.store.time.time', return_value=1700001000):
                    self.assertEqual(client.get('/api/reference/json', headers=headers).status_code, 423)
                with patch('challenge.store.time.time', return_value=1700001030):
                    batch = client.get('/api/reference/json', headers=headers).json()
                self.assertIn('to red', batch['target_prompt'])
                result = client.post('/api/reference/submissions', headers=headers,
                                     json=dict(batch_id=batch['batch_id'], edited_image=batch['image']))
                self.assertEqual(result.status_code, 201, result.text)
                self.assertEqual(result.json()['scoring_status'], 'pending')
                self.assertIsNone(result.json()['metrics'])
                self.assertTrue(result.json()['excel_updated'])
                self.assertEqual(client.get('/api/reference/leaderboard', headers=headers).json(), [])
                duplicate = client.post('/api/reference/submissions', headers=headers,
                                        json=dict(batch_id=batch['batch_id'], edited_image=batch['image']))
                self.assertEqual(duplicate.status_code, 400)
                book = load_workbook(root / 'results.xlsx')
                try:
                    page = book['Reference uploads']
                    headings = [cell.value for cell in page[1]]
                    self.assertIsNone(page.cell(2, headings.index('psnr') + 1).value)
                    self.assertIsNone(page.cell(2, headings.index('mse') + 1).value)
                    self.assertEqual(book['Team totals'].max_row, 2)
                finally:
                    book.close()
            store = Store(root)
            try:
                stored = store.db.execute('SELECT metrics,source,mask FROM reference_batches').fetchone()
                self.assertIsNone(stored['metrics'])
                self.assertIsNotNone(stored['source'])
                self.assertIsNone(stored['mask'])
            finally:
                store.db.close()

    def test_demo_and_existing_schema_migration(self):
        with tempfile.TemporaryDirectory() as directory:
            root = initialize_reference(Path(directory) / 'demo')
            store = Store(root)
            self.assertNotIn('ground_truth_mask', store.reference_task)
            self.assertNotIn('mask_url', store.public_task(store.reference_task))
            store.db.execute('ALTER TABLE reference_batches RENAME TO old_batches')
            store.db.execute('CREATE TABLE reference_batches AS SELECT id,team,window_end,source_hash,width,height,created,received,edited FROM old_batches')
            store.db.execute('DROP TABLE old_batches')
            store.db.commit()
            store.db.close()
            store = Store(root)
            try:
                columns = {r['name'] for r in store.db.execute('PRAGMA table_info(reference_batches)')}
                self.assertTrue({'metrics', 'mask', 'source'} <= columns)
            finally:
                store.db.close()
            with self.assertRaisesRegex(ValueError, 'refusing to overwrite'):
                initialize_reference(root)
