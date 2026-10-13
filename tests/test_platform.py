import base64
import io
import json
import tempfile
import threading
import unittest
from unittest.mock import patch, Mock
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import Request, urlopen

import numpy as np
from PIL import Image
from openpyxl import load_workbook

from challenge.baseline import edit, infer
from challenge.demo import initialize
from challenge.credentials import add_team
from challenge.metrics import evaluate, localization, score
from challenge.runner import collect, execute
from challenge.server import Server
from challenge.store import Invalid, Store, image_bytes


def encoded(image):
    return base64.b64encode(image_bytes(image)).decode()


class Fixture(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.config = initialize(self.root)
        self.store = Store(self.root)

    def tearDown(self):
        self.store.db.close()
        self.tmp.cleanup()

    def payload(self, task_id='validation_001', run='run_1', unchanged=False):
        task = self.store.tasks[task_id]
        source = self.store.load(task, 'source_image', 'RGB')
        edited, mask, metadata = edit(source, task['source_prompt'], task['target_prompt'])
        metadata['run_id'] = run
        return dict(task_id=task_id, student_id='team_demo', edited_image=encoded(source if unchanged else edited),
                    predicted_mask=encoded(mask), metadata=metadata)


class ReferenceTests(Fixture):
    def test_edited_upload_is_bound_to_team_batch_and_persists(self):
        raw = image_bytes(Image.new('RGB', (16, 12), 'blue'))
        batch_id = self.store.issue_reference_batch('team_demo', raw, 16, 12, 1060)
        self.assertEqual(batch_id, self.store.issue_reference_batch('team_demo', raw, 16, 12, 1060))
        payload = dict(batch_id=batch_id, edited_image=encoded(Image.new('RGB', (16, 12), 'red')))
        self.store.config['teams']['other_team'] = 'x' * 32
        for team, admin in (('other_team', False), ('admin', True)):
            with self.assertRaises(PermissionError):
                self.store.submit_reference(payload, team, admin)
        with self.assertRaises(PermissionError):
            self.store.submit_reference(dict(payload, batch_id='unknown'), 'team_demo')
        for image in (Image.new('RGB', (8, 8)), Image.new('L', (16, 12))):
            with self.assertRaises(Invalid):
                self.store.submit_reference(dict(payload, edited_image=encoded(image)), 'team_demo')
        other = Store(self.root)
        try:
            result = other.submit_reference(payload, 'team_demo')
            self.assertEqual(result['status'], 'accepted')
            self.assertEqual(result['scoring_status'], 'pending')
            self.assertTrue(result['excel_updated'])
            with self.assertRaises(Invalid):
                other.submit_reference(payload, 'team_demo')
            row = other.db.execute('SELECT edited FROM reference_batches WHERE id=?', (batch_id,)).fetchone()
            self.assertEqual(row['edited'], image_bytes(Image.new('RGB', (16, 12), 'red')))
            book = load_workbook(self.root / 'results.xlsx')
            try:
                self.assertEqual(book['Reference uploads']['B2'].value, batch_id)
                self.assertEqual(book['Reference uploads']['A2'].value, 'team_demo')
            finally:
                book.close()
        finally:
            other.db.close()

    def test_json_download_then_upload_route(self):
        from challenge.server import Handler
        handler = Handler.__new__(Handler)
        handler.server = Mock(store=self.store)
        handler.reply = Mock()
        handler.path = '/api/reference/json'
        handler.headers = {'Authorization': 'Bearer ' + self.config['teams']['team_demo']}
        with patch('challenge.store.time.time', return_value=1000):
            handler.route(False)
        with patch('challenge.store.time.time', return_value=1030):
            handler.route(False)
        downloaded = handler.reply.call_args.args[1]
        payload = json.dumps({'batch_id': downloaded['batch_id'], 'edited_image': downloaded['image']}).encode()
        handler.headers.update({'Content-Type': 'application/json', 'Content-Length': str(len(payload))})
        handler.rfile = io.BytesIO(payload)
        handler.path = '/api/reference/submissions'
        handler.route(True)
        self.assertEqual(handler.reply.call_args.args[0], 201)
        self.assertEqual(handler.reply.call_args.args[1]['batch_id'], downloaded['batch_id'])

    def test_json_bundle_auth_timing_image_and_private_data(self):
        from challenge.server import Handler, REFERENCE
        handler = Handler.__new__(Handler)
        handler.server = Mock(store=self.store)
        handler.reply = Mock()
        handler.path = '/api/reference/json'
        for headers in ({}, {'Authorization': 'Bearer invalid'}):
            handler.headers = headers
            handler.route(False)
            self.assertEqual(handler.reply.call_args.args[0], 401)
        handler.headers = {'Authorization': 'Bearer ' + self.config['teams']['team_demo']}
        for now, status in ((1000, 423), (1030, 200), (1060, 423)):
            with patch('challenge.store.time.time', return_value=now):
                handler.route(False)
            code, payload = handler.reply.call_args.args
            self.assertEqual(code, status)
            if status == 423:
                self.assertNotIn('image', payload)
                continue
            decoded = base64.b64decode(payload['image'], validate=True)
            self.assertEqual(decoded, REFERENCE.read_bytes())
            with Image.open(io.BytesIO(decoded)) as image:
                self.assertEqual(image.size, (payload['image_width'], payload['image_height']))
            public = json.dumps(payload)
            for secret in (self.config['admin_token'], *self.config['teams'].values()):
                self.assertNotIn(secret, public)
            self.assertNotIn('ground_truth_mask', payload)

    def test_boundaries_repeat_and_restart_does_not_reset_wait(self):
        for now, available, remaining in ((1000, False, 30), (1029, False, 1),
                                          (1030, True, 30), (1059, True, 1),
                                          (1060, False, 30), (1090, True, 30)):
            with patch('challenge.store.time.time', return_value=now):
                window = self.store.reference_window('team_demo')
            self.assertEqual(window['available'], available)
            self.assertEqual(window['seconds_remaining'], remaining)
        other = Store(self.root)
        try:
            with patch('challenge.store.time.time', return_value=1095):
                self.assertEqual(other.reference_window('team_demo')['seconds_remaining'], 25)
                self.assertFalse(other.reference_window('another_team')['available'])
        finally:
            other.db.close()

    def test_reference_routes_require_token_and_enforce_window(self):
        from challenge.server import Handler, REFERENCE
        handler = Handler.__new__(Handler)
        handler.server = Mock(store=self.store)
        handler.reply = Mock()
        handler.path = '/api/reference/image'
        handler.headers = {}
        handler.route(False)
        self.assertEqual(handler.reply.call_args.args[0], 401)
        handler.headers = {'Authorization': 'Bearer wrong'}
        handler.route(False)
        self.assertEqual(handler.reply.call_args.args[0], 401)
        handler.headers = {'Authorization': 'Bearer ' + self.config['teams']['team_demo']}
        with patch('challenge.store.time.time', return_value=1000):
            handler.route(False)
            self.assertEqual(handler.reply.call_args.args[0], 423)
            handler.path = '/api/reference'
            handler.route(False)
            self.assertEqual(handler.reply.call_args.args[1]['seconds_remaining'], 30)
        handler.path = '/api/reference/image'
        with patch('challenge.store.time.time', return_value=1030):
            handler.route(False)
            self.assertEqual(handler.reply.call_args.args, (200, REFERENCE.read_bytes(), 'image/png'))
        with patch('challenge.store.time.time', return_value=1060):
            handler.route(False)
            self.assertEqual(handler.reply.call_args.args[0], 423)
        handler.path = '/data/image.png'
        with self.assertRaises(KeyError):
            handler.route(False)


class CredentialTests(Fixture):
    def test_issue_unique_tokens_and_preserve_configuration(self):
        first = add_team(self.root, 'team_01')
        second = add_team(self.root, 'team_02')
        self.assertNotEqual(first, second)
        self.assertGreaterEqual(len(first), 24)
        path = self.root / 'config.json'
        config = json.loads(path.read_text())
        self.assertEqual(config['admin_token'], self.config['admin_token'])
        self.assertEqual(config['teams']['team_demo'], self.config['teams']['team_demo'])
        self.assertEqual(path.stat().st_mode & 0o777, 0o600)
        before = path.read_bytes()
        for team in ('team_01', '../bad', '', 'a' * 81):
            with self.assertRaises(ValueError):
                add_team(self.root, team)
        self.assertEqual(path.read_bytes(), before)
        reloaded = Store(self.root)
        try:
            self.assertEqual(reloaded.authenticate(first), ('team_01', False))
            self.assertEqual(reloaded.authenticate(second), ('team_02', False))
            self.assertEqual(reloaded.authenticate('invalid-token'), (None, False))
        finally:
            reloaded.db.close()


class ExcelTests(Fixture):
    def rows(self, name):
        book = load_workbook(self.root / 'results.xlsx')
        try:
            values = list(book[name].values)
            return [dict(zip(values[0], row)) for row in values[1:]]
        finally:
            book.close()

    def test_metrics_rankings_selection_and_restart(self):
        result = self.store.submit(self.payload(), 'team_demo')
        self.assertTrue(result['excel_updated'])
        record = self.rows('Submissions')[0]
        self.assertEqual(record['Team'], 'team_demo')
        self.assertEqual(record['score'], 1)
        self.assertEqual(record['dice'], 1)
        self.assertIsNone(record['efficiency'])
        quality = next(r for r in self.rows('Team rankings') if r['Mode'] == 'quality')
        self.assertEqual(quality['Score'], .5)  # One of two tasks is missing.
        self.store.submit(self.payload(run='worse', unchanged=True), 'team_demo')
        self.store.select('team_demo', 'validation', 'run_1')
        self.assertEqual(self.rows('Team rankings')[0]['Selected run'], 'run_1')
        (self.root / 'results.xlsx').unlink()
        other = Store(self.root)
        other.db.close()
        self.assertEqual(len(self.rows('Submissions')), 2)

    def test_export_failure_preserves_records_and_previous_workbook(self):
        before = (self.root / 'results.xlsx').read_bytes()
        with patch('challenge.store.Path.replace', side_effect=PermissionError('Workbook locked')):
            with self.assertLogs(level='ERROR'):
                result = self.store.submit(self.payload(), 'team_demo')
        self.assertEqual(result['status'], 'accepted')
        self.assertFalse(result['excel_updated'])
        self.assertEqual(len(self.store.submissions('team_demo')), 1)
        self.assertEqual((self.root / 'results.xlsx').read_bytes(), before)
        self.assertEqual(list(self.root.glob('*.xlsx')), [self.root / 'results.xlsx'])
        self.store.export_excel()
        self.assertEqual(len(self.rows('Submissions')), 1)

    def test_formula_like_text_is_literal_and_null_scores_stay_blank(self):
        self.store.tasks['validation_001']['task_type'] = '=1+1'
        self.store.tasks['validation_001'].pop('edit_rubric')
        self.store.submit(self.payload(), 'team_demo')
        book = load_workbook(self.root / 'results.xlsx')
        try:
            self.assertEqual(book['Submissions']['F2'].value, '=1+1')
            self.assertEqual(book['Submissions']['F2'].data_type, 's')
            self.assertIsNone(self.rows('Submissions')[0]['score'])
        finally:
            book.close()


class SharedImageTests(Fixture):
    def test_upload_receive_expiration_and_restart(self):
        payload = {'filename': 'source.png', 'image': encoded(Image.new('RGB', (16, 12), 'red'))}
        with patch('challenge.store.time.time', return_value=1000):
            result = self.store.upload_image(payload, admin=True)
            self.assertEqual(result['expires_at'], 1300)
            self.assertEqual(len(self.store.shared_images()), 1)
        other = Store(self.root)
        try:
            with patch('challenge.store.time.time', return_value=1299):
                with Image.open(io.BytesIO(other.shared_image(result['image_id']))) as image:
                    self.assertEqual(image.size, (16, 12))
            with patch('challenge.store.time.time', return_value=1300):
                with self.assertRaises(KeyError):
                    other.shared_image(result['image_id'])
                self.assertEqual(other.shared_images(), [])
        finally:
            other.db.close()

    def test_authorization_validation_and_storage_limit(self):
        payload = {'image': encoded(Image.new('RGB', (16, 12)))}
        with self.assertRaises(PermissionError):
            self.store.upload_image(payload)
        for filename in ('../private.png', '', 'a\\b.png', 'bad\n.png'):
            with self.assertRaises(Invalid):
                self.store.upload_image(dict(payload, filename=filename), admin=True)
        with self.assertRaises(Invalid):
            self.store.upload_image({'image': 'bad-base64'}, admin=True)
        with patch('challenge.store.uuid.uuid4') as identifier:
            for i in range(100):
                identifier.return_value.hex = str(i)
                self.store.upload_image(payload, admin=True)
            with self.assertRaises(Invalid):
                self.store.upload_image(payload, admin=True)

    def test_rgb_jpeg_is_normalized_to_png(self):
        stream = io.BytesIO()
        Image.new('RGB', (16, 12)).save(stream, format='JPEG')
        result = self.store.upload_image({'image': base64.b64encode(stream.getvalue()).decode()}, admin=True)
        self.assertTrue(self.store.shared_image(result['image_id']).startswith(b'\x89PNG'))


class MetricTests(Fixture):
    def test_perfect_edit_and_unchanged_image_are_distinguished(self):
        result = self.store.submit(self.payload(), 'team_demo')['metrics']
        for key in ('iou', 'dice', 'precision', 'recall', 'boundary_f1', 'editing', 'preservation', 'score'):
            self.assertAlmostEqual(result[key], 1.)
        result = self.store.submit(self.payload(run='unchanged', unchanged=True), 'team_demo')['metrics']
        self.assertEqual(result['editing'], 0.)
        self.assertEqual(result['preservation'], 1.)
        self.assertLess(result['score'], 1.)

    def test_empty_prediction_and_boundary_tolerance(self):
        truth = self.store.load(self.store.tasks['validation_001'], 'ground_truth_mask', 'L')
        empty = Image.new('L', truth.size)
        result = localization(empty, truth)
        for key in ('iou', 'dice', 'precision', 'recall', 'boundary_f1'):
            self.assertEqual(result[key], 0)
        with self.assertRaises(ValueError):
            localization(empty, empty)
        self.assertEqual(localization(truth, truth, tolerance=0)['boundary_f1'], 1)

    def test_background_uses_reference_region_not_prediction(self):
        payload = self.payload()
        payload['predicted_mask'] = encoded(Image.new('L', (256, 192), 255))
        edited = Image.new('RGB', (256, 192), (0, 0, 0))
        payload['edited_image'] = encoded(edited)
        result = self.store.submit(payload, 'team_demo')['metrics']
        self.assertGreater(result['background_mae'], .5)
        self.assertLess(result['iou'], .2)

    def test_unconfigured_edit_metric_does_not_create_fake_score(self):
        task = dict(self.store.tasks['validation_001'])
        task.pop('edit_rubric')
        source = self.store.load(task, 'source_image', 'RGB')
        truth = self.store.load(task, 'ground_truth_mask', 'L')
        result = evaluate(source, source, truth, truth, truth, task, self.config)
        self.assertIsNone(result['editing'])
        self.assertIsNone(score(result, self.config))

    def test_full_image_allowed_region_rejected(self):
        task = self.store.tasks['validation_001']
        source = self.store.load(task, 'source_image', 'RGB')
        truth = self.store.load(task, 'ground_truth_mask', 'L')
        with self.assertRaises(ValueError):
            evaluate(source, source, truth, truth, Image.new('L', source.size, 255), task, self.config)


class SubmissionTests(Fixture):
    def test_duplicate_does_not_destroy_previous_submission(self):
        payload = self.payload()
        self.store.submit(payload, 'team_demo')
        with self.assertRaises(Invalid):
            self.store.submit(payload, 'team_demo')
        self.assertEqual(len(self.store.submissions('team_demo')), 1)

    def test_mask_mode_dimensions_and_base64_validation(self):
        for encoded_mask in (encoded(Image.new('RGB', (256, 192))), encoded(Image.new('L', (1, 1))), 'not base64'):
            with self.subTest(encoded_mask=encoded_mask[:20]):
                payload = self.payload()
                payload['predicted_mask'] = encoded_mask
                with self.assertRaises(Invalid):
                    self.store.submit(payload, 'team_demo')
        self.assertEqual(self.store.submissions('team_demo'), [])

    def test_reported_times_and_identity_cannot_be_forged(self):
        for value in (-1, float('nan'), float('inf'), True, 'fast'):
            payload = self.payload()
            payload['metadata']['editing_time_ms'] = value
            with self.assertRaises(Invalid):
                self.store.submit(payload, 'team_demo')
        payload = self.payload()
        payload['student_id'] = 'other'
        with self.assertRaises(PermissionError):
            self.store.submit(payload, 'team_demo')
        payload = self.payload()
        payload['measured_total_time_ms'] = 1
        with self.assertRaises(Invalid):
            self.store.submit(payload, 'team_demo')
        with self.assertRaises(PermissionError):
            self.store.submit(self.payload(), 'team_demo', measured=True)

    def test_test_submissions_do_not_expose_metrics(self):
        result = self.store.submit(self.payload('test_001'), 'team_demo')
        self.assertNotIn('metrics', result)
        self.assertNotIn('metrics', self.store.submissions('team_demo')[0])
        self.assertIn('metrics', self.store.submissions('admin', True)[0])
        public = self.store.public_task(self.store.tasks['test_001'])
        for key in ('ground_truth_mask', 'edit_rubric', 'mask_url', 'source_image'):
            self.assertNotIn(key, public)

    def test_artifacts_are_owned(self):
        result = self.store.submit(self.payload(), 'team_demo')
        with self.assertRaises(PermissionError):
            self.store.artifact(result['submission_id'], 'edited', 'other', False)
        self.assertTrue(self.store.artifact(result['submission_id'], 'edited', 'admin', True).startswith(b'\x89PNG'))

    def test_results_survive_restart_and_changed_benchmark_is_rejected(self):
        self.store.submit(self.payload(), 'team_demo')
        other = Store(self.root)
        self.assertEqual(len(other.submissions('team_demo')), 1)
        other.db.close()
        config = dict(self.config)
        config['weights']['editing'] = .5
        (self.root / 'config.json').write_text(json.dumps(config))
        with self.assertRaises(Invalid):
            Store(self.root)

    def test_dataset_path_escape_rejected(self):
        with self.assertRaises(Invalid):
            self.store.path('../secret')

    def test_reference_mask_must_be_binary(self):
        Image.new('L', (256, 192), 128).save(self.root / 'masks/validation_001.png')
        with self.assertRaises(Invalid):
            Store(self.root)

    def test_limit_enforced_across_runs(self):
        self.store.config['max_team_submissions'] = 1
        self.store.submit(self.payload(), 'team_demo')
        with self.assertRaises(Invalid):
            self.store.submit(self.payload(run='new_run'), 'team_demo')


class LeaderboardTests(Fixture):
    def test_entirely_failed_registered_runs_score_zero(self):
        self.store.register_run('team_demo', 'validation', 'failed', measured=True)
        self.store.select('team_demo', 'validation', 'failed')
        for measured in (False, True):
            result = self.store.leaderboard('validation', measured)[0]
            self.assertEqual(result['score'], 0)
            self.assertEqual(result['completed'], 0)

    def test_task_types_are_weighted_equally(self):
        self.store.tasks['validation_002']['task_type'] = 'fine_grained'
        extra = dict(self.store.tasks['validation_001'], task_id='validation_003')
        self.store.tasks['validation_003'] = extra
        self.store.submit(self.payload(), 'team_demo')
        self.store.submit(self.payload('validation_003'), 'team_demo')
        self.assertAlmostEqual(self.store.leaderboard('validation')[0]['score'], .5)

    def test_missing_tasks_count_as_zero_and_selected_run_is_used(self):
        self.store.submit(self.payload(), 'team_demo')
        row = self.store.leaderboard('validation')[0]
        self.assertAlmostEqual(row['score'], .5)
        self.assertEqual(row['completed'], 1)
        self.store.submit(self.payload('validation_002'), 'team_demo')
        self.assertAlmostEqual(self.store.leaderboard('validation')[0]['score'], 1)
        self.store.select('team_demo', 'validation', 'run_1')
        self.store.submit(self.payload(run='newest', unchanged=True), 'team_demo')
        self.assertEqual(self.store.leaderboard('validation')[0]['run'], 'run_1')

    def test_uploads_and_measured_execution_do_not_mix(self):
        self.store.submit(self.payload(), 'team_demo')
        self.assertIsNone(self.store.leaderboard('validation', True)[0]['score'])
        payload = self.payload('validation_002')
        payload.update(measured_total_time_ms=10000, hardware_profile=self.config['hardware_profile'])
        result = self.store.submit(payload, 'admin', admin=True, measured=True)
        self.assertEqual(result['metrics']['efficiency'], .5)
        row = self.store.leaderboard('validation', True)[0]
        self.assertEqual(row['completed'], 1)
        self.assertAlmostEqual(row['score'], .4625)
        quality = self.store.leaderboard('validation')[0]
        self.assertAlmostEqual(quality['score'], 1.)

    def test_hardware_profile_must_match(self):
        payload = self.payload()
        payload.update(measured_total_time_ms=1, hardware_profile='unapproved_gpu')
        with self.assertRaises(Invalid):
            self.store.submit(payload, 'admin', admin=True, measured=True)


class HTTPTests(Fixture):
    def setUp(self):
        super().setUp()
        self.server = Server(('127.0.0.1', 0), self.store)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.url = f'http://127.0.0.1:{self.server.server_port}'

    def tearDown(self):
        self.server.shutdown()
        self.thread.join()
        self.server.server_close()
        super().tearDown()

    def request(self, path, payload=None, admin=False, token=True):
        headers = {'Content-Type': 'application/json'}
        if token:
            headers['Authorization'] = 'Bearer ' + (self.config['admin_token'] if admin else self.config['teams']['team_demo'])
        req = Request(self.url + path, data=json.dumps(payload).encode() if payload is not None else None, headers=headers)
        return urlopen(req, timeout=5)

    def test_source_image_upload_and_download_api(self):
        payload = {'filename': 'lecture.png', 'image': encoded(Image.new('RGB', (16, 12)))}
        with self.assertRaises(HTTPError) as error:
            self.request('/api/images', payload)
        self.assertEqual(error.exception.code, 403)
        with self.request('/api/images', payload, admin=True) as response:
            result = json.load(response)
            self.assertEqual(response.status, 201)
            self.assertEqual(result['duration_seconds'], 300)
        with self.request('/api/images') as response:
            self.assertEqual(json.load(response)[0]['image_id'], result['image_id'])
        with self.request(result['image_url']) as response:
            self.assertEqual(response.headers['Content-Type'], 'image/png')
        with self.assertRaises(HTTPError) as error:
            self.request(result['image_url'], token=False)
        self.assertEqual(error.exception.code, 401)
        with patch('challenge.store.time.time', return_value=result['expires_at']):
            with self.assertRaises(HTTPError) as error:
                self.request(result['image_url'])
            self.assertEqual(error.exception.code, 404)

    def test_dashboard_and_authenticated_task_api(self):
        with self.request('/', token=False) as response:
            self.assertEqual(response.status, 200)
            self.assertIn(b'Semantic Editing Lab', response.read())
            self.assertIn("frame-ancestors 'none'", response.headers['Content-Security-Policy'])
        with self.assertRaises(HTTPError) as error:
            self.request('/api/tasks', token=False)
        self.assertEqual(error.exception.code, 401)
        with self.request('/api/tasks?split=validation') as response:
            self.assertEqual(len(json.load(response)), 2)

    def test_new_team_token_accesses_tasks_but_not_private_masks(self):
        token = add_team(self.root, 'team_01')
        reloaded = Store(self.root)
        self.server.store = reloaded
        try:
            headers = {'Authorization': 'Bearer ' + token}
            with urlopen(Request(self.url + '/api/tasks?split=validation', headers=headers)) as response:
                task = json.load(response)[0]
            with urlopen(Request(self.url + task['image_url'], headers=headers)) as response:
                self.assertTrue(response.read().startswith(b'\x89PNG'))
            with self.assertRaises(HTTPError) as error:
                urlopen(Request(self.url + '/api/tasks/test_001/mask', headers=headers))
            self.assertEqual(error.exception.code, 403)
        finally:
            self.server.store = self.store
            reloaded.db.close()

    def test_private_test_mask_and_no_arbitrary_files(self):
        for path, code in (('/api/tasks/test_001/mask', 403), ('/config.json', 404), ('/.challenge/config.json', 404)):
            with self.assertRaises(HTTPError) as error:
                self.request(path)
            self.assertEqual(error.exception.code, code)
        with self.request('/api/tasks/test_001/mask', admin=True) as response:
            self.assertEqual(response.headers['Content-Type'], 'image/png')

    def test_submit_and_query_end_to_end(self):
        with self.request('/api/submissions', self.payload()) as response:
            result = json.load(response)
            self.assertEqual(response.status, 201)
            self.assertEqual(result['metrics']['score'], 1)
        with self.request('/api/submissions') as response:
            self.assertEqual(len(json.load(response)), 1)
        with self.request('/api/submissions/' + str(result['submission_id']) + '/edited') as response:
            self.assertTrue(response.read().startswith(b'\x89PNG'))

    def test_excel_download_is_organizer_only_and_current(self):
        for token, status in ((False, 401), (True, 403)):
            with self.assertRaises(HTTPError) as error:
                self.request('/api/results.xlsx', token=token)
            self.assertEqual(error.exception.code, status)
        with self.request('/api/submissions', self.payload(task_id='test_001')) as response:
            self.assertNotIn('metrics', json.load(response))
        with self.request('/api/results.xlsx', admin=True) as response:
            self.assertEqual(response.headers['Content-Type'],
                             'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            book = load_workbook(io.BytesIO(response.read()))
        try:
            self.assertEqual(book['Submissions'].max_row, 2)
            self.assertEqual(book['Submissions']['D2'].value, 'test_001')
        finally:
            book.close()

    def test_nonfinite_json_is_rejected_and_admin_token_not_exposed(self):
        payload = self.payload()
        payload['metadata']['editing_time_ms'] = float('nan')
        with self.assertRaises(HTTPError) as error:
            self.request('/api/submissions', payload)
        self.assertEqual(error.exception.code, 400)
        with self.request('/api/config', admin=True) as response:
            public = response.read().decode()
            self.assertNotIn(self.config['admin_token'], public)
            self.assertNotIn(self.config['teams']['team_demo'], public)


class UploadDeadlineTests(unittest.TestCase):
    def test_total_receive_deadline_cannot_be_extended_by_chunks(self):
        from challenge.server import Handler
        handler = Handler.__new__(Handler)
        handler.connection = Mock()
        handler.rfile = Mock()
        handler.rfile.read1.side_effect = [b'a', b'b']
        with patch('challenge.server.time.monotonic', side_effect=[0, 0, 59, 61]):
            with self.assertRaises(TimeoutError):
                handler.read_upload(3)
        self.assertEqual(handler.rfile.read1.call_count, 2)
        self.assertEqual(handler.connection.settimeout.call_args.args, (15,))

    def test_success_and_incomplete_upload(self):
        from challenge.server import Handler
        handler = Handler.__new__(Handler)
        handler.connection = Mock()
        handler.rfile = Mock()
        handler.rfile.read1.side_effect = [b'a', b'b']
        with patch('challenge.server.time.monotonic', return_value=0):
            self.assertEqual(handler.read_upload(2), b'ab')
        handler.rfile.read1.side_effect = [b'']
        with self.assertRaises(Invalid):
            handler.read_upload(1)


class RunnerTests(Fixture):
    def test_inference_contract_produces_valid_artifacts(self):
        task = self.store.tasks['validation_001']
        (self.root / 'task.json').write_text(json.dumps(task))
        output = self.root / 'output'
        infer(self.root / 'task.json', output)
        files, metadata = collect(output)
        self.assertIn('localization_time_ms', metadata)
        with Image.open(io.BytesIO(files['mask.png'])) as mask:
            self.assertEqual(mask.mode, 'L')
        with Image.open(io.BytesIO(files['edited.png'])) as edited:
            self.assertEqual(edited.size, (256, 192))

    def test_symlink_outputs_and_timeout_are_rejected(self):
        output = self.root / 'output'
        output.mkdir()
        (output / 'edited.png').symlink_to(self.root / 'config.json')
        with self.assertRaises(ValueError):
            collect(output)
        import sys
        with self.assertRaises(ValueError):
            execute([sys.executable, '-c', 'import time; time.sleep(2)'], .05)


if __name__ == '__main__':
    unittest.main()
