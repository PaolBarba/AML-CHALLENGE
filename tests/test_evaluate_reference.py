import base64
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock

from PIL import Image
from openpyxl import load_workbook

from challenge.demo import initialize
from challenge.evaluate_reference import evaluate_batches
from challenge.reference_metrics import pixel_metrics
from challenge.store import Store, image_bytes


class EvaluateReferenceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        config = initialize(root)
        config['reference_task'] = 'validation_001'
        (root / 'config.json').write_text(json.dumps(config))
        self.store = Store(root)
        self.source = self.store.load(self.store.reference_task, 'source_image', 'RGB')
        self.edited = Image.new('RGB', self.source.size, 'red')
        self.batches = []
        for end in (1060, 1120):
            batch = self.store.issue_reference_batch('team_demo', image_bytes(self.source), *self.source.size, end)
            self.store.submit_reference({'batch_id': batch, 'edited_image': base64.b64encode(image_bytes(self.edited)).decode()}, 'team_demo')
            self.batches.append(batch)
        self.waiting = self.store.issue_reference_batch('team_demo', image_bytes(self.source), *self.source.size, 1180)

    def tearDown(self):
        self.store.db.close()
        self.tmp.cleanup()

    def metrics(self, similarity):
        return dict(pixel_metrics(self.source, self.edited), clip_similarity=similarity, evaluator_version='test-clip')

    def test_evaluate_all_accepted_batches_updates_excel_and_ranking(self):
        evaluator = Mock()
        evaluator.evaluate.side_effect = [self.metrics(.7), self.metrics(.8)]
        self.assertEqual(evaluate_batches(self.store, evaluator), (2, True))
        evaluator.evaluate.assert_called_with(unittest.mock.ANY, unittest.mock.ANY, self.store.reference_task)
        board = self.store.reference_leaderboard()
        self.assertEqual(board[0]['batch_id'], self.batches[1])
        self.assertEqual(board[0]['score'], .8)
        waiting = self.store.db.execute('SELECT metrics FROM reference_batches WHERE id=?', (self.waiting,)).fetchone()
        self.assertIsNone(waiting['metrics'])
        book = load_workbook(self.store.root / 'results.xlsx')
        try:
            headers = [cell.value for cell in book['Reference uploads'][1]]
            from challenge.reference_metrics import METRICS
            self.assertEqual(headers, ['Team ID', 'Batch ID', *METRICS])
            self.assertEqual(book.sheetnames, ['Reference uploads', 'Team totals'])
            totals = book['Team totals']
            self.assertEqual(totals['A2'].value, 'team_demo')
            columns = [cell.value for cell in totals[1]]
            self.assertAlmostEqual(totals.cell(2, columns.index('clip_similarity') + 1).value, 1.5)
            self.assertIsNone(totals.cell(2, columns.index('lpips') + 1).value)
        finally:
            book.close()

    def test_failure_does_not_overwrite_any_previous_scores(self):
        previous = [r['metrics'] for r in self.store.db.execute('SELECT metrics FROM reference_batches ORDER BY id')]
        evaluator = Mock()
        evaluator.evaluate.side_effect = [self.metrics(.9), RuntimeError('model failed')]
        with self.assertRaisesRegex(RuntimeError, 'model failed'):
            evaluate_batches(self.store, evaluator)
        self.assertEqual(previous, [r['metrics'] for r in self.store.db.execute('SELECT metrics FROM reference_batches ORDER BY id')])

    def test_missing_clip_does_not_replace_pixel_metrics(self):
        evaluator = Mock()
        evaluator.evaluate.return_value = self.metrics(None)
        with self.assertRaisesRegex(ValueError, 'local CLIP model'):
            evaluate_batches(self.store, evaluator)
        self.assertEqual(self.store.reference_leaderboard(), [])

    def test_api_dataset_without_reference_task_uses_explicit_prompt(self):
        self.store.reference_task = None
        evaluator = Mock()
        evaluator.evaluate.side_effect = [self.metrics(.7), self.metrics(.8)]
        with self.assertRaisesRegex(ValueError, '--target-prompt'):
            evaluate_batches(self.store, evaluator)
        self.assertEqual(evaluate_batches(self.store, evaluator, 'Change the backpack to red'), (2, True))
        self.assertEqual(evaluator.evaluate.call_args.args[2]['target_prompt'], 'Change the backpack to red')

    def test_reference_evaluation_does_not_need_other_benchmark_assets(self):
        fingerprint = self.store.db.execute('SELECT fingerprint FROM benchmark').fetchone()[0]
        self.store.path(self.store.tasks['train_001']['source_image']).unlink()
        evaluator = Mock()
        evaluator.evaluate.side_effect = [self.metrics(.7), self.metrics(.8)]
        offline = Store(self.store.root, reference_only=True)
        try:
            self.assertEqual(evaluate_batches(offline, evaluator), (2, True))
            self.assertEqual(offline.db.execute('SELECT fingerprint FROM benchmark').fetchone()[0], fingerprint)
        finally:
            offline.db.close()
